import os, re, json, uuid, time, shutil, hmac, difflib
from urllib.parse import urlparse
from flask import Flask, request, jsonify, send_from_directory, Response
from werkzeug.utils import secure_filename
import requests

BASE = os.path.dirname(os.path.abspath(__file__))
UPLOADS = os.path.join(BASE, 'uploads')
os.makedirs(UPLOADS, exist_ok=True)
VERIFY_TOKEN = os.getenv('VERIFY_TOKEN', 'fastai123')
READABLE = {'txt', 'csv', 'tsv', 'json', 'log', 'md', 'xlsx', 'xlsm', 'pdf'}
MESSAGES, CONNECTED = [], set()

app = Flask(__name__, static_folder=None)
app.config['MAX_CONTENT_LENGTH'] = 300 * 1024 * 1024

def ext_of(n): return n.rsplit('.', 1)[-1].lower() if '.' in n else ''
def safe_name(orig):
    e = ext_of(orig); n = secure_filename(orig)
    return n if n.endswith('.' + e) and len(n) > len(e) + 1 else 'file.' + e
def info(fid):
    p = os.path.join(UPLOADS, fid)
    return {'id': fid, 'name': fid.split('_', 1)[1], 'size': os.path.getsize(p), 'type': ext_of(fid)}

@app.route('/')
def home(): return send_from_directory(BASE, 'index.html')
@app.route('/style.css')
def css(): return send_from_directory(BASE, 'style.css')
@app.route('/app.js')
def js(): return send_from_directory(BASE, 'app.js')

# ---------- AUTH ----------
@app.before_request
def guard():
    pw = os.getenv('APP_PASSWORD')
    if not pw or request.path.startswith('/api/webhook/'): return
    a = request.authorization
    if not a or not hmac.compare_digest(a.password or '', pw):
        return Response('Login required', 401, {'WWW-Authenticate': 'Basic realm="Fast AI"'})

# ---------- PROVIDER REGISTRY ----------
with open(os.path.join(BASE, 'providers.json'), encoding='utf-8') as _f:
    REG = json.load(_f)['providers']

def _host(u): return urlparse(u if '://' in u else 'https://' + u).netloc.lower()
def clean_base(b):
    b = (b or '').strip().rstrip('/')
    for suf in ('/chat/completions', '/completions', '/models'):
        if b.endswith(suf): b = b[:-len(suf)]
    return b.rstrip('/')

def resolve(key, base, model):
    base, key, model, p = clean_base(base), (key or '').strip(), (model or '').strip(), None
    if base:
        if '://' not in base: base = 'https://' + base
        p = next((x for x in REG if _host(x['base']) == _host(base)), None)
        if p and urlparse(base).path in ('', '/'): base = p['base']
    if not p and key:
        hits = [(len(pre), x) for x in REG for pre in x.get('keyPrefix', []) if key.startswith(pre)]
        if hits: p = max(hits, key=lambda h: h[0])[1]
    if p and not base: base = p['base']
    if not base: base = 'https://api.openai.com/v1'
    if p and not model: model = p['models'][0]
    return p, base, model

def chat(base, key, body, timeout=90):
    hdr = {'Authorization': f'Bearer {key or "none"}', 'Content-Type': 'application/json'}
    b2 = {k: v for k, v in body.items() if k != 'response_format'}
    b3 = {k: v for k, v in b2.items() if k not in ('temperature', 'max_tokens')}
    b3['max_completion_tokens'] = body.get('max_tokens', 3000)
    for b in (body, b2, b3):
        r = requests.post(base + '/chat/completions', headers=hdr, json=b, timeout=timeout)
        if r.status_code != 400: break
    return r

def friendly(r):
    c = r.status_code
    if c in (401, 403): return 'API key galat hai ya permission nahi (%d)' % c
    if c == 404: return 'URL ya model galat hai (404): ' + r.text[:160]
    if c == 429: return 'Limit/credits khatam (429): ' + r.text[:120]
    return f'API error {c}: {r.text[:200]}'

@app.get('/api/providers')
def providers(): return jsonify(providers=[{k: p[k] for k in ('id', 'name', 'base', 'models')} for p in REG])

@app.post('/api/resolve')
def api_resolve():
    d = request.json or {}
    p, base, model = resolve(d.get('api_key'), d.get('api_base'), d.get('api_model'))
    return jsonify(provider=p and {'id': p['id'], 'name': p['name']}, base=base, model=model, models=p['models'] if p else [])

@app.post('/api/test')
def test_connection():
    d = request.json or {}
    key = (d.get('api_key') or '').strip()
    p, base, model = resolve(key, d.get('api_base'), d.get('api_model'))
    name = p['name'] if p else _host(base)
    if not key and not (p and p.get('local')): return jsonify(ok=False, error='API key daalo')
    t, models = time.time(), []
    try:
        r = requests.get(base + '/models', headers={'Authorization': f'Bearer {key or "none"}'}, timeout=20)
        if r.status_code in (401, 403): return jsonify(ok=False, provider=name, base=base, error=friendly(r))
        if r.status_code == 200:
            js = r.json()
            models = sorted({(m.get('id') or m.get('name') or '').replace('models/', '') for m in (js.get('data') or js.get('models') or []) if isinstance(m, dict)} - {''})
        if not model and models: model = models[0]
        r = chat(base, key, {'model': model, 'messages': [{'role': 'user', 'content': 'hi'}], 'max_tokens': 5}, 40)
    except Exception as e:
        return jsonify(ok=False, provider=name, base=base, error='Connect nahi hua: ' + str(e)[:160])
    ms = int((time.time() - t) * 1000)
    if r.status_code != 200:
        return jsonify(ok=False, provider=name, base=base, model=model, models=models, error=friendly(r),
                       suggestions=difflib.get_close_matches(model, models, 4, 0.3))
    return jsonify(ok=True, provider=name, base=base, model=model, models=models or (p['models'] if p else []), ms=ms)

# ---------- FILES ----------
@app.post('/api/upload')
def upload():
    out = []
    for f in request.files.getlist('files'):
        orig = os.path.basename(f.filename or '')
        if ext_of(orig) not in READABLE: continue
        fid = uuid.uuid4().hex[:8] + '_' + safe_name(orig)
        f.save(os.path.join(UPLOADS, fid)); out.append(info(fid))
    if not out: return jsonify(error='Koi supported file nahi mili (txt, csv, xlsx, pdf, json, md, log)'), 400
    return jsonify(files=out)

@app.get('/api/files')
def files():
    return jsonify(files=[info(f) for f in sorted(os.listdir(UPLOADS)) if '_' in f])

@app.delete('/api/delete/<fid>')
def delete(fid):
    p = os.path.join(UPLOADS, secure_filename(fid))
    if os.path.isfile(p): os.remove(p)
    return jsonify(ok=True)

@app.post('/api/connect-folder')
def connect_folder():
    path = (request.json or {}).get('path', '').strip().strip('"')
    if not os.path.isdir(path): return jsonify(error='Folder nahi mila. Poora path daalo.'), 400
    added = []
    for root, _, names in os.walk(path):
        for n in names:
            src = os.path.join(root, n)
            if ext_of(n) in READABLE and len(added) < 800 and os.path.getsize(src) < 40 * 1024 * 1024:
                fid = uuid.uuid4().hex[:8] + '_' + safe_name(n)
                shutil.copy2(src, os.path.join(UPLOADS, fid)); added.append(info(fid))
    return jsonify(files=added)

def read_text(path, limit):
    e = ext_of(path)
    try:
        if e in ('txt', 'csv', 'tsv', 'json', 'log', 'md'):
            return open(path, encoding='utf-8', errors='ignore').read(limit)
        if e in ('xlsx', 'xlsm'):
            from openpyxl import load_workbook
            wb, out = load_workbook(path, read_only=True, data_only=True), []
            for ws in wb.worksheets[:8]:
                out.append('## Sheet: ' + ws.title)
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 500: break
                    out.append(','.join('' if c is None else str(c) for c in row))
            return '\n'.join(out)[:limit]
        if e == 'pdf':
            from pypdf import PdfReader
            return '\n'.join((p.extract_text() or '') for p in PdfReader(path).pages[:25])[:limit]
    except Exception:
        pass
    return ''

# ---------- AI (improved prompt for conversational + accurate charts) ----------
SYSTEM = """You are Fast AI — a smart, friendly business analyst for the company owner. Talk like a trusted friend (Hinglish / Hindi / English matching the question).

STRICT RULES:
1. Use ONLY the exact numbers and data present in the given files. Never invent or assume numbers.
2. If data is missing, clearly say so.
3. Always return ONLY valid JSON (no markdown, no extra text).

JSON FORMAT:
{
  "spoken": "1-2 short friendly sentences that you will speak out loud. Sound natural like talking to a friend.",
  "answer": "Detailed but easy explanation in short paragraphs. Use simple language. When talking about loss/profit mention exact numbers from files.",
  "kpis": [{"label":"short label","value":"₹ 25,000 or whatever exact","tone":"profit|loss|neutral"}],
  "charts": [
    {
      "title": "clear chart title",
      "type": "bar|line|doughnut",
      "tone": "profit|loss|neutral",
      "unit": "₹ or ₹ Lakh or %",
      "labels": ["2020","2021",...],
      "data": [numbers matching labels exactly from files]
    }
  ],
  "causes": ["only if loss/problem question — likely reasons based on data"],
  "solutions": ["only if loss/problem — concrete next steps"]
}

CHART RULES (very important):
- Create a SEPARATE chart for every major point (yearly totals, monthly trend, comparison, category split…).
- Use real numbers from files. If you find “25000 loss” or “Rahul” related data, highlight it clearly.
- Prefer bar for comparisons, line for trends over years/months, doughnut for splits.
- Make charts beautiful and useful so the owner can understand at a glance.

TONE:
- Speak like a helpful younger brother / smart friend.
- Example spoken: “Bhai, 5 saal pehle 2021 mein ₹25,000 ka loss dikha tha files mein. Main chart bana ke dikhata hoon.”
- Keep answer clear, short paragraphs, no jargon.
"""

def demo():
    return {
        'spoken': 'Ye demo jawab hai. Asli jawab ke liye API key daalo.',
        'answer': '⚠️ DEMO MODE — API key nahi mili, isliye ye sample data hai, aapki files ka nahi. Setup screen me API key daalo.',
        'kpis': [
            {'label': 'Profit (sample)', 'value': '₹ 17 Lakh', 'tone': 'profit'},
            {'label': 'Loss (sample)', 'value': '₹ 6 Lakh', 'tone': 'loss'}
        ],
        'charts': [{
            'title': 'Profit Trend (sample)',
            'type': 'bar',
            'tone': 'profit',
            'unit': '₹ Lakh',
            'labels': ['2021', '2022', '2023', '2024', '2025'],
            'data': [8, 11, 14, 17, 20]
        }],
        'causes': [],
        'solutions': []
    }

@app.post('/api/analyze')
def analyze():
    d = request.json or {}
    key = (d.get('api_key') or '').strip()
    if not key and not (resolve('', d.get('api_base'), '')[0] or {}).get('local'):
        return jsonify(demo())
    p, base, model = resolve(key, d.get('api_base'), d.get('api_model') or 'gpt-4o-mini')
    ctx, budget = [], 55000
    for f in sorted(os.listdir(UPLOADS)):
        t = read_text(os.path.join(UPLOADS, f), 9000)
        if t.strip() and budget > 0:
            ctx.append(f'=== FILE: {f.split("_", 1)[-1]} ===\n{t}')
            budget -= len(t)
    body = {
        'model': model,
        'temperature': 0.25,
        'max_tokens': 3500,
        'response_format': {'type': 'json_object'},
        'messages': [
            {'role': 'system', 'content': SYSTEM},
            {'role': 'user', 'content': 'FILES:\n' + ('\n'.join(ctx) or 'No readable files') + '\n\nQUESTION: ' + d.get('query', '')}
        ]
    }
    try:
        r = chat(base, key, body)
        if r.status_code != 200:
            return jsonify(error=friendly(r)), 400
        txt = r.json()['choices'][0]['message']['content']
        return jsonify(json.loads(txt[txt.find('{'): txt.rfind('}') + 1]))
    except Exception as e:
        return jsonify(error=str(e)[:300]), 400

# ---------- APPS / MESSAGES ----------
def add_msg(app_name, sender, text):
    MESSAGES.append({'ts': time.time(), 'app': app_name, 'from': sender, 'text': text})
    del MESSAGES[:-200]

@app.post('/api/connect')
def connect():
    d = request.json or {}
    (CONNECTED.add if d.get('on') else CONNECTED.discard)(d.get('app'))
    return jsonify(ok=True, webhook=f"{request.host_url}api/webhook/{d.get('app')}")

@app.route('/api/webhook/<name>', methods=['GET', 'POST'])
def webhook(name):
    if request.method == 'GET':
        if request.args.get('hub.verify_token') == VERIFY_TOKEN:
            return request.args.get('hub.challenge', '')
        return 'forbidden', 403
    d = request.get_json(silent=True) or {}
    if name == 'whatsapp' and 'entry' in d:
        for e in d['entry']:
            for ch in e.get('changes', []):
                v = ch.get('value', {})
                names = {c.get('wa_id'): c.get('profile', {}).get('name') for c in v.get('contacts', [])}
                for m in v.get('messages', []):
                    add_msg('whatsapp', names.get(m.get('from')) or m.get('from'), (m.get('text') or {}).get('body', '[media]'))
    else:
        add_msg(name, d.get('from', 'Unknown'), d.get('text', ''))
    return jsonify(ok=True)

@app.get('/api/messages')
def messages():
    since = float(request.args.get('since', 0))
    return jsonify(messages=[m for m in MESSAGES if m['ts'] > since], now=time.time())

@app.post('/api/messages/test')
def test_msg():
    add_msg((request.json or {}).get('app', 'whatsapp'), 'Rahul Sharma', 'Sir, kal ki meeting 4 baje shift kar do?')
    return jsonify(ok=True)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.getenv('PORT', 5000)), debug=False)
